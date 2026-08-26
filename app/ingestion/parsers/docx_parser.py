"""DOCX document parser using python-docx."""

from __future__ import annotations

import io

from docx import Document as DocxDocument

from app.ingestion.parsers.base import DocumentParser
from app.schemas.common import ParsedDocument, RawDocument


class DocxParser(DocumentParser):
    """Extract paragraphs and heading-based sections from DOCX files."""

    def supports(self, content_type: str, filename: str) -> bool:
        return filename.lower().endswith(".docx") or content_type in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }

    def parse(self, raw: RawDocument) -> ParsedDocument:
        document = DocxDocument(io.BytesIO(raw.content))
        sections: list[dict[str, object]] = []
        current_section: dict[str, object] | None = None
        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            paragraphs.append(text)
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            if "heading" in style_name:
                if current_section:
                    sections.append(current_section)
                current_section = {
                    "title": text,
                    "content": "",
                    "level": style_name,
                }
            elif current_section is not None:
                existing = str(current_section.get("content", ""))
                current_section["content"] = f"{existing}\n{text}".strip()
            else:
                current_section = {"title": "Introduction", "content": text}

        if current_section:
            sections.append(current_section)

        core_props = document.core_properties
        title = core_props.title or raw.reference.filename.rsplit(".", 1)[0]
        metadata = {
            "author": core_props.author,
            "subject": core_props.subject,
            "paragraph_count": len(paragraphs),
        }

        return ParsedDocument(
            text="\n\n".join(paragraphs),
            title=title,
            metadata=metadata,
            sections=sections,
        )
